"""Build/validate a review-only paper using python-docx; never replace the source.

Run from the repository root with python-docx installed. No database or network use.
"""
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/defense/BPA_MODULES_UPDATED_OCT_22_2026.docx'
TARGET = SOURCE.with_name('BPA_MODULES_UPDATED_OCT_22_2026_A9_REVIEW.docx')
AMENDMENT = SOURCE.with_name('BPA_OCT22_A9_SECTION_AMENDMENT.md')
REPORT = SOURCE.with_name('BPA_A9_REVIEW_VALIDATION.md')


def append_note(cell, text):
    reference = cell.paragraphs[-1]
    paragraph = cell.add_paragraph(style=reference.style)
    if reference._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(reference._p.pPr))
    run = paragraph.add_run(text)
    if reference.runs and reference.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(reference.runs[0]._r.rPr))


source_bytes = SOURCE.read_bytes()
before = Document(SOURCE)
document = Document(SOURCE)
amendment = AMENDMENT.read_text(encoding='utf-8')
blocks = list(re.finditer(r'^## Module (\d+): ([^\n]+)\n(.*?)(?=^## |\Z)', amendment, re.M | re.S))
assert [int(m[1]) for m in blocks] == list(range(1, 26))
assert len(before.tables) == 328
changes = []
for block in blocks:
    number = int(block[1])
    tables = document.tables[1 + (number - 1) * 13:1 + number * 13]
    assert len(tables) == 13
    assert tables[0].cell(1, 1).text == before.tables[0].cell(number, 1).text
    rows = re.findall(r'^\| (\d+)\. [^|]+ \| (.+) \|$', block[3], re.M)
    if rows:
        for section, text in rows:
            section = int(section)
            assert section in (2, 3, 4, 5, 7, 9, 10, 11, 13)
            table = tables[section - 1]
            # Overview has Purpose then Process Description; append only to its last cell.
            append_note(table.rows[-1].cells[-1], text)
            changes.append((number, section))
    else:
        text = re.search(r'In section 13 append: "(.*?)"', block[3], re.S)
        assert text, number
        append_note(tables[12].rows[-1].cells[-1], text[1])
        changes.append((number, 13))

append_note(document.tables[1 + 7 * 13 + 12].rows[-1].cells[-1],
    'Snapshot update, 2026-10-30: B6 commit 4c993244416ec9b0530e1d16acfb6c9e471fac9c '
    'reports two guarded synthetic API tests (33 assertions), including HTTP 201 with persisted '
    'Visit/Queue identity and idempotency, plus separate controlled rendered success/error checks. '
    'This supersedes pending-report wording from the earlier amendment; live browser-to-server '
    'and shared integration acceptance are still unverified. Evidence: '
    'docs/STF12_B6_WALKIN_ACCEPTANCE.md on that isolated commit, not on shared integration.')

notice = document.paragraphs[3].insert_paragraph_before(
    'A9 REVIEW REQUIRED — revised candidate, 2026-10-30 (Asia/Manila). '
    'Original October 22 paper preserved separately. Structural validation passed; '
    'Word pagination, table flow and visual formatting require human review before approval.'
)
notice.runs[0].bold = True
anchor = next(p for p in document.paragraphs if p.text.startswith('Module 1  |'))
anchor.insert_paragraph_before(
    'Development evidence clarification, 2026-10-30: declared dependencies are Laravel ^13.17 '
    'and React ^19.1.1. Shared integration is 010a386f147bed285c58821a0a2cc800c393982e. '
    'A8 and B3/B4/B5 are isolated. B6 report 4c993244416ec9b0530e1d16acfb6c9e471fac9c '
    'provides isolated synthetic API success and rendered fixture evidence; live browser-to-server '
    'acceptance and shared integration remain unverified. Earlier pending-B6 wording below '
    'belongs to the approved amendment snapshot. Original proposed-process and interview text '
    'is preserved; dated qualifications limit its implementation claims. See '
    'ACADEMIC_DOCUMENT_RECONCILIATION.md and FEATURE_ACCEPTANCE_MATRIX.md. '
    'The historical 20 Implemented / 5 Partial inventory is not reclassified.'
)
document.save(TARGET)

# Reopen the saved artifact and prove preservation, not merely successful serialization.
after = Document(TARGET)
assert SOURCE.read_bytes() == source_bytes
assert len(after.tables) == len(before.tables) == 328
assert len(after.sections) == len(before.sections) == 1
assert [p.text for p in before.paragraphs if p.text.startswith('Module ')] == [
    p.text for p in after.paragraphs if p.text.startswith('Module ')
]
expected = {(n, s) for n, s in changes}
for i, (old, new) in enumerate(zip(before.tables, after.tables)):
    assert (len(old.rows), len(old.columns)) == (len(new.rows), len(new.columns))
    assert etree.tostring(old._tbl.tblPr) == etree.tostring(new._tbl.tblPr)
    assert etree.tostring(old._tbl.tblGrid) == etree.tostring(new._tbl.tblGrid)
    for r, (old_row, new_row) in enumerate(zip(old.rows, new.rows)):
        for c, (old_cell, new_cell) in enumerate(zip(old_row.cells, new_row.cells)):
            assert etree.tostring(old_cell._tc.tcPr) == etree.tostring(new_cell._tc.tcPr)
            original_paragraphs = old_cell.paragraphs
            revised_paragraphs = new_cell.paragraphs
            assert len(revised_paragraphs) >= len(original_paragraphs)
            for a, b in zip(original_paragraphs, revised_paragraphs):
                assert etree.tostring(a._p) == etree.tostring(b._p), (i, r, c)
            if len(revised_paragraphs) > len(original_paragraphs):
                assert i > 0
                assert ((i - 1) // 13 + 1, (i - 1) % 13 + 1) in expected
            assert new_cell.text.strip() or not old_cell.text.strip()
assert etree.tostring(before.sections[0]._sectPr) == etree.tostring(after.sections[0]._sectPr)
changed_parts = []
with ZipFile(SOURCE) as old, ZipFile(TARGET) as new:
    assert new.testzip() is None
    assert {n for n in old.namelist() if not n.endswith('/')} == {
        n for n in new.namelist() if not n.endswith('/')
    }  # ZIP directory entries are not document parts and python-docx omits them.
    for name in new.namelist():
        content = new.read(name)
        if name.endswith(('.xml', '.rels')):
            etree.fromstring(content)
        if content != old.read(name):
            changed_parts.append(name)
    assert changed_parts == ['word/document.xml'], changed_parts
    # This also covers all header/footer, style, theme, relationship and media parts.
    old_paragraphs = etree.fromstring(old.read('word/document.xml')).xpath('//w:p', namespaces={
        'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    })
    new_paragraphs = etree.fromstring(new.read('word/document.xml')).xpath('//w:p', namespaces={
        'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    })
    iterator = iter(new_paragraphs)
    for paragraph in old_paragraphs:
        assert any(etree.tostring(p) == etree.tostring(paragraph) for p in iterator)

REPORT.write_text(
    '# A9 Word candidate validation\n\n'
    '2026-10-30, Asia/Manila. Generated by [build_a9_review.py](build_a9_review.py).\n\n'
    '- [Original](BPA_MODULES_UPDATED_OCT_22_2026.docx): unchanged.\n'
    '- [Revised candidate](BPA_MODULES_UPDATED_OCT_22_2026_A9_REVIEW.docx): **REVIEW REQUIRED**.\n'
    '- ZIP CRC, XML/relationship parsing and python-docx reopen: passed.\n'
    '- 25 module titles/owners and assignment rows, 328 tables, every row/column count, '
    'table/cell properties and grid widths, and one layout section preserved.\n'
    '- Every original paragraph XML preserved in sequence, including all interview/proposed text; '
    'no original academic paragraph removed or rewritten.\n'
    '- Only word/document.xml changed; ZIP directory entries are omitted by python-docx. All actual '
    'package parts, including styles, headers/footers, '
    'relationships and media, remain byte-identical.\n'
    f'- {len(changes)} expected module/section content cells received qualification paragraphs; '
    'one further B6 snapshot paragraph in the already-qualified M8 section 13 cell and '
    'two body notices added. No new table.\n'
    f'- Original SHA-256: `{sha256(source_bytes).hexdigest()}`.\n'
    f'- Candidate SHA-256: `{sha256(TARGET.read_bytes()).hexdigest()}`.\n\n'
    'Visual attempt: Microsoft Word COM activation failed with REGDB_E_CLASSNOTREG '
    '(0x80040154); WINWORD.EXE and LibreOffice/soffice were not found. No rendered-layout '
    'acceptance is claimed and the original was not replaced.\n\n'
    'Human review: open candidate in Word or LibreOffice; check all module starts, long '
    'qualified cells and adjoining page breaks for clipped text, split rows, blank pages and '
    'overflow; inspect cover notice, assignment table, headers/footers, signatures, margins '
    'and fonts. Export PDF and inspect all pages before accepting it as the final paper.\n',
    encoding='utf-8'
)
print(f'PASS: 25 modules, 328 tables, {len(changes)} expected cell additions; original unchanged')
